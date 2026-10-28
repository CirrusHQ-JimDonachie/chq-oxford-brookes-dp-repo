"""Turn the Archibus room-data workbook into the CSV the reference/rooms config reads.

The workbook carries a title block above its header row, line breaks inside
some header cells and a long run of empty rows, none of which the engine's CSV
reader can skip. This keeps the first sheet's header row and every non-empty
row after it, with header whitespace collapsed to single spaces.

    uv run --no-project --with openpyxl python mapping-configs/rooms_xlsx_to_csv.py \
        "RS NHHB-Room Data-Sept 2026.xlsx" nhhb-rooms.csv
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

HEADER_MARKER = "Room Code"


def main(source: str, target: str) -> None:
    """Write the workbook's room rows to a CSV the engine can read."""
    sheet = load_workbook(source, read_only=True, data_only=True).worksheets[0]
    rows = iter(sheet.iter_rows(values_only=True))
    for row in rows:
        header = [" ".join(str(cell or "").split()) for cell in row]
        if HEADER_MARKER in header:
            break
    else:
        sys.exit(f"No header row containing {HEADER_MARKER!r} in {source}")

    written = 0
    with Path(target).open("w", encoding="utf-8", newline="") as out:
        writer = csv.writer(out)
        writer.writerow(header)
        for row in rows:
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            writer.writerow(["" if cell is None else str(cell).strip() for cell in row])
            written += 1
    sys.stdout.write(f"{written} rooms written to {target}\n")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
